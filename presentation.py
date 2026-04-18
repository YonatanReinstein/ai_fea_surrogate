"""Generate project presentation PowerPoint."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ── Color palette ──
BG_DARK   = RGBColor(0x1B, 0x1B, 0x2F)
BG_MID    = RGBColor(0x22, 0x22, 0x3A)
ACCENT    = RGBColor(0x4F, 0xC3, 0xF7)
ACCENT2   = RGBColor(0x66, 0xBB, 0x6A)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT     = RGBColor(0xCC, 0xCC, 0xDD)
ORANGE    = RGBColor(0xFF, 0xA7, 0x26)
RED_SOFT  = RGBColor(0xEF, 0x53, 0x50)

def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_textbox(slide, left, top, width, height, text, font_size=18,
                color=WHITE, bold=False, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return tf

def add_bullet_slide(slide, left, top, width, height, items, font_size=18,
                     color=WHITE, bullet_color=ACCENT, spacing=Pt(8)):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.space_after = spacing
        # bullet char
        run_b = p.add_run()
        run_b.text = "\u25B8 "
        run_b.font.size = Pt(font_size)
        run_b.font.color.rgb = bullet_color
        run_b.font.name = "Calibri"
        # text
        run_t = p.add_run()
        run_t.text = item
        run_t.font.size = Pt(font_size)
        run_t.font.color.rgb = color
        run_t.font.name = "Calibri"
    return tf

def add_code_box(slide, left, top, width, height, text, font_size=13):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                                   Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0x15, 0x15, 0x25)
    shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.15)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = ACCENT
    p.font.name = "Consolas"
    return tf

def section_bar(slide, text=""):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                                   Inches(13.333), Inches(0.08))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()

def slide_number(slide, num, total):
    add_textbox(slide, 12.3, 7.0, 1, 0.4, f"{num}/{total}", font_size=11,
                color=LIGHT, alignment=PP_ALIGN.RIGHT)

total_slides = 14

# ═══════════════════════════════════════════════════════════════
# SLIDE 1 — Title
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])  # blank
set_slide_bg(sl, BG_DARK)
section_bar(sl)

add_textbox(sl, 1, 1.8, 11, 1.2,
            "AI-Based Surrogate Modeling\nfor Finite Element Analysis",
            font_size=40, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

add_textbox(sl, 1, 3.6, 11, 0.6,
            "Replacing Expensive FEA Simulations with Graph Neural Networks\nfor Rapid Parametric Design Optimization",
            font_size=20, color=LIGHT, alignment=PP_ALIGN.CENTER)

add_textbox(sl, 1, 5.2, 11, 0.5,
            "Yonatan Reinstein",
            font_size=22, color=ACCENT, bold=True, alignment=PP_ALIGN.CENTER)

slide_number(sl, 1, total_slides)

# ═══════════════════════════════════════════════════════════════
# SLIDE 2 — Motivation / Problem Statement
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(sl, BG_DARK)
section_bar(sl)
add_textbox(sl, 0.8, 0.4, 10, 0.7, "The Problem: FEA is a Bottleneck", font_size=32, color=ACCENT, bold=True)

add_bullet_slide(sl, 0.8, 1.4, 5.5, 4.5, [
    "FEA solves PDEs on meshes to predict stress, displacement, etc.",
    "Each ANSYS simulation takes seconds to minutes per design",
    "Design optimization requires evaluating thousands of candidates",
    "Brute-force search over parametric dimensions is infeasible",
    "Goal: train a fast surrogate model that approximates FEA\n  with orders-of-magnitude speedup",
], font_size=18)

# right column — conceptual comparison
add_textbox(sl, 7.2, 1.4, 5, 0.5, "Traditional vs. Surrogate", font_size=20, color=ORANGE, bold=True)
add_code_box(sl, 7.2, 2.1, 5.3, 1.6,
    "Traditional:  dims \u2192 CAD \u2192 Mesh \u2192 ANSYS \u2192 stress\n"
    "              ~30 sec per design \u00d7 10,000 designs\n"
    "              = ~83 hours",
    font_size=14)
add_code_box(sl, 7.2, 4.0, 5.3, 1.6,
    "Surrogate:    dims \u2192 Neural Network \u2192 stress\n"
    "              ~5 ms per design \u00d7 10,000 designs\n"
    "              = ~50 seconds",
    font_size=14)

slide_number(sl, 2, total_slides)

# ═══════════════════════════════════════════════════════════════
# SLIDE 3 — System Overview / Pipeline
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(sl, BG_DARK)
section_bar(sl)
add_textbox(sl, 0.8, 0.4, 10, 0.7, "System Overview: Four-Stage Pipeline", font_size=32, color=ACCENT, bold=True)

stages = [
    ("1. Data Generation", "Parametric CAD\n\u2192 IRIT Mesh\n\u2192 ANSYS FEA\n\u2192 Graph Dataset", ACCENT),
    ("2. Model Training", "GNN / MLP\nSurrogate learns\nFEA mapping", ACCENT2),
    ("3. Optimization", "Genetic Algorithm\nuses surrogate\nfor fast evaluation", ORANGE),
    ("4. Evaluation", "Compare candidates\nVisualize results\nVerify with FEA", RED_SOFT),
]
for i, (title, desc, color) in enumerate(stages):
    x = 0.6 + i * 3.15
    shape = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(1.5), Inches(2.8), Inches(3.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG_MID
    shape.line.color.rgb = color
    shape.line.width = Pt(2)

    add_textbox(sl, x + 0.15, 1.7, 2.5, 0.5, title, font_size=20, color=color, bold=True, alignment=PP_ALIGN.CENTER)
    add_textbox(sl, x + 0.15, 2.4, 2.5, 2.0, desc, font_size=16, color=LIGHT, alignment=PP_ALIGN.CENTER)

    if i < 3:
        add_textbox(sl, x + 2.75, 2.7, 0.5, 0.5, "\u25B6", font_size=24, color=LIGHT, alignment=PP_ALIGN.CENTER)

add_textbox(sl, 0.8, 5.2, 11, 0.8,
    "The surrogate model is trained offline on FEA data, then used online during optimization\n"
    "to evaluate thousands of design candidates in seconds instead of hours.",
    font_size=16, color=LIGHT)

slide_number(sl, 3, total_slides)

# ═══════════════════════════════════════════════════════════════
# SLIDE 4 — Data Generation: From Geometry to Graphs
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(sl, BG_DARK)
section_bar(sl)
add_textbox(sl, 0.8, 0.4, 10, 0.7, "Stage 1: Data Generation", font_size=32, color=ACCENT, bold=True)

add_bullet_slide(sl, 0.8, 1.3, 5.5, 5, [
    "IRIT parametric CAD model defines geometry from dimensions\n  (e.g., 36 parameters d1\u2013d36 for the \"arm\" geometry)",
    "Random sampling of dimension values within bounds\n  (uniform distribution, seeded for reproducibility)",
    "IRIT generates hex/tet mesh for each parameter set",
    "ANSYS MAPDL runs FEA simulation:\n  \u2022 Material: Steel (E = 200 GPa, \u03BD = 0.3)\n  \u2022 BCs: anchored base + applied tip load (1 MN)\n  \u2022 Output: von Mises stress, displacement at every node",
    "Parallel execution: 4 MAPDL instances with retry logic",
    "~1000\u20133000 samples generated per geometry",
], font_size=17)

add_textbox(sl, 7.2, 1.3, 5, 0.5, "Pipeline per Sample", font_size=20, color=ORANGE, bold=True)
add_code_box(sl, 7.2, 2.0, 5.3, 4.5,
    "dims = {d1: 0.41, d2: 0.73, ..., d36: 2.20}\n"
    "          \u2193\n"
    "IRIT CAD model  \u2192  volume = 17.87 m\u00b3\n"
    "          \u2193\n"
    "Hex mesh (10\u00d710\u00d720)  \u2192  2000+ nodes\n"
    "          \u2193\n"
    "ANSYS MAPDL solve\n"
    "          \u2193\n"
    "Per-node: [Ux, Uy, Uz, \u03c3_vm]\n"
    "Global:   max_stress = 7.36e7 Pa\n"
    "          max_disp   = 0.0012 m",
    font_size=14)

slide_number(sl, 4, total_slides)

# ═══════════════════════════════════════════════════════════════
# SLIDE 5 — What is a Neural Network? (basics for CG professor)
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(sl, BG_DARK)
section_bar(sl)
add_textbox(sl, 0.8, 0.4, 10, 0.7, "Background: Neural Networks in 60 Seconds", font_size=32, color=ACCENT, bold=True)

add_bullet_slide(sl, 0.8, 1.3, 5.5, 5, [
    "A neural network is a parametric function f(x; \u03b8) \u2192 y",
    "Composed of layers: each applies a linear transform\n  followed by a non-linear activation (e.g., ReLU)",
    "Training = finding \u03b8 that minimizes a loss function\n  L(\u03b8) = \u2211 || f(x\u1d62; \u03b8) - y\u1d62 ||\u00b2  (mean squared error)",
    "Optimization via gradient descent:\n  \u03b8 \u2190 \u03b8 - \u03b1 \u00b7 \u2207L(\u03b8)  (backpropagation computes gradients)",
    "Universal approximation: with enough parameters,\n  can approximate any continuous function",
    "Key advantage: once trained, inference is O(ms)\n  vs. O(seconds) for FEA",
], font_size=17)

add_textbox(sl, 7.2, 1.3, 5, 0.5, "MLP: Multi-Layer Perceptron", font_size=20, color=ORANGE, bold=True)
add_code_box(sl, 7.2, 2.0, 5.3, 2.0,
    "Input (36 dims)\n"
    "   \u2193  Linear(36, 128) + ReLU\n"
    "Hidden (128)\n"
    "   \u2193  Linear(128, 128) + ReLU\n"
    "Hidden (128)\n"
    "   \u2193  Linear(128, 3)\n"
    "Output: [volume, stress, displacement]",
    font_size=14)

add_textbox(sl, 7.2, 4.3, 5.3, 1.5,
    "The MLP takes dimension values as a flat vector and directly predicts "
    "global outputs. Simple and fast, but ignores the spatial/mesh structure.",
    font_size=15, color=LIGHT)

slide_number(sl, 5, total_slides)

# ═══════════════════════════════════════════════════════════════
# SLIDE 6 — From Meshes to Graphs
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(sl, BG_DARK)
section_bar(sl)
add_textbox(sl, 0.8, 0.4, 10, 0.7, "Key Insight: FEA Meshes Are Graphs", font_size=32, color=ACCENT, bold=True)

add_bullet_slide(sl, 0.8, 1.3, 6.0, 2.5, [
    "A FEA mesh is naturally a graph: nodes = vertices, elements define edges",
    "Each node carries features: position (x,y,z), forces (Fx,Fy,Fz), boundary flag",
    "Each node has labels from FEA: displacement (Ux,Uy,Uz) and stress (\u03c3_vm)",
    "This is exactly the structure Graph Neural Networks are designed to process",
], font_size=18)

add_textbox(sl, 0.8, 3.8, 5.5, 0.5, "Graph Data Object (PyTorch Geometric)", font_size=20, color=ORANGE, bold=True)
add_code_box(sl, 0.8, 4.4, 5.5, 2.5,
    "Data(\n"
    "  x          = [N, 7]   # node features\n"
    "               [x, y, z, Fx, Fy, Fz, anchored]\n"
    "  edge_index = [2, E]   # bidirectional edges\n"
    "  volume     = scalar   # graph-level label\n"
    "  max_stress = scalar   # graph-level label\n"
    "  dims       = [1, 36]  # parametric dimensions\n"
    ")",
    font_size=14)

add_textbox(sl, 7.2, 3.8, 5.3, 0.5, "Why Graphs Instead of Flat Vectors?", font_size=20, color=ORANGE, bold=True)
add_bullet_slide(sl, 7.2, 4.4, 5.3, 2.5, [
    "Spatial locality: stress depends on neighboring nodes",
    "Variable mesh sizes: graphs handle any number of nodes",
    "Topology-aware: captures connectivity, not just coordinates",
    "Physically meaningful: mirrors how FEA itself works",
], font_size=17, color=LIGHT)

slide_number(sl, 6, total_slides)

# ═══════════════════════════════════════════════════════════════
# SLIDE 7 — GNN Architecture (the core AI)
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(sl, BG_DARK)
section_bar(sl)
add_textbox(sl, 0.8, 0.4, 10, 0.7, "The GNN Surrogate Model", font_size=32, color=ACCENT, bold=True)

add_textbox(sl, 0.8, 1.3, 5.5, 0.5, "Architecture Overview", font_size=22, color=ORANGE, bold=True)
add_code_box(sl, 0.8, 1.9, 5.5, 4.5,
    "Node features [N, 7]\n"
    "        \u2193\n"
    "  Encoder MLP:  7 \u2192 128 \u2192 128\n"
    "        \u2193\n"
    "  6\u00d7 EdgeConv layers (with residual connections)\n"
    "    Each layer:\n"
    "      h\u1d62 = h\u1d62 + EdgeConv(h\u1d62, edges)\n"
    "      EdgeConv aggregates neighbor info via:\n"
    "        MLP([h\u1d62 || h\u2c7c], 128, 128), aggr = max\n"
    "        \u2193\n"
    "  Head MLP:  128 \u2192 128 \u2192 1  (per-node prediction)\n"
    "        \u2193\n"
    "  Global Max Pool: aggregate all nodes \u2192 1 scalar\n"
    "        \u2193\n"
    "  Output: predicted max stress",
    font_size=14)

add_textbox(sl, 7.2, 1.3, 5.3, 0.5, "Key Concepts Explained", font_size=22, color=ORANGE, bold=True)
add_bullet_slide(sl, 7.2, 1.9, 5.3, 5, [
    "EdgeConv: for each node, concatenate its features\n  with each neighbor's features, apply MLP, take max\n  \u2192 learns edge-level interactions",
    "Residual connections: h = h + conv(h)\n  \u2192 preserves information across layers\n  \u2192 stabilizes training of deep networks",
    "Message passing: each layer lets nodes \"communicate\"\n  with neighbors \u2192 after 6 layers, information has\n  propagated 6 hops across the mesh",
    "Global max pool: reduces variable-size node\n  predictions to a single graph-level output\n  (analogous to taking max stress over all nodes)",
], font_size=16)

slide_number(sl, 7, total_slides)

# ═══════════════════════════════════════════════════════════════
# SLIDE 8 — EdgeConv deep dive (visual / intuitive)
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(sl, BG_DARK)
section_bar(sl)
add_textbox(sl, 0.8, 0.4, 10, 0.7, "EdgeConv: Learning from Neighbors", font_size=32, color=ACCENT, bold=True)

add_textbox(sl, 0.8, 1.3, 11, 0.8,
    "EdgeConv is the core building block. Think of it as: each mesh node looks at its neighbors,\n"
    "computes a learned function of the pair, and aggregates the results.",
    font_size=18, color=LIGHT)

add_textbox(sl, 0.8, 2.4, 5.5, 0.5, "Mathematical Definition", font_size=20, color=ORANGE, bold=True)
add_code_box(sl, 0.8, 3.0, 5.5, 2.5,
    "For node i with neighbors N(i):\n\n"
    "  h\u1d62' = MAX   { MLP( [h\u1d62 || h\u2c7c] ) }\n"
    "       j \u2208 N(i)\n\n"
    "  [h\u1d62 || h\u2c7c] = concatenation of features\n"
    "  MLP         = learnable non-linear function\n"
    "  MAX         = element-wise maximum over neighbors",
    font_size=15)

add_textbox(sl, 7.2, 2.4, 5.3, 0.5, "Physical Analogy", font_size=20, color=ORANGE, bold=True)
add_bullet_slide(sl, 7.2, 3.0, 5.3, 3.5, [
    "Similar to how FEA assembles element stiffness:\n  each element contributes to its nodes' equations",
    "EdgeConv: each edge contributes a learned\n  message to its endpoint nodes",
    "Max aggregation picks the dominant influence\n  (like max stress being determined by the most\n  stressed neighboring element)",
    "After 6 layers: each node's representation encodes\n  information from a 6-hop neighborhood in the mesh",
], font_size=16)

slide_number(sl, 8, total_slides)

# ═══════════════════════════════════════════════════════════════
# SLIDE 9 — Training Process
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(sl, BG_DARK)
section_bar(sl)
add_textbox(sl, 0.8, 0.4, 10, 0.7, "Stage 2: Training the Surrogate", font_size=32, color=ACCENT, bold=True)

add_textbox(sl, 0.8, 1.3, 5.5, 0.5, "Training Setup", font_size=22, color=ORANGE, bold=True)
add_bullet_slide(sl, 0.8, 1.9, 5.5, 3.5, [
    "Dataset: ~3000 FEA simulations stored as graphs",
    "Split: 80% training / 20% validation (random)",
    "Loss function: Mean Squared Error (MSE)\n  L = || predicted_stress - true_stress ||\u00b2",
    "Optimizer: AdamW (lr = 0.002, weight_decay = 1e-4)\n  \u2192 Adam with decoupled weight regularization",
    "Scheduler: StepLR (step=20, \u03b3=0.9)\n  \u2192 learning rate decays every 20 epochs",
    "Normalization: features & targets scaled to ~unit range\n  (forces /1e6, stress /1e6 \u2192 work in MPa & MN)",
], font_size=16)

add_textbox(sl, 7.2, 1.3, 5.3, 0.5, "Training Convergence", font_size=22, color=ORANGE, bold=True)
add_code_box(sl, 7.2, 1.9, 5.3, 2.5,
    "Epoch  1:  train_loss = 0.980  val_loss = 0.95\n"
    "Epoch 10:  train_loss = 0.087  val_loss = 0.09\n"
    "Epoch 25:  train_loss = 0.031  val_loss = 0.03\n"
    "Epoch 50:  train_loss = 0.019  val_loss = 0.02\n\n"
    "  \u2192 ~98% reduction in prediction error\n"
    "  \u2192 val \u2248 train: model generalizes well (no overfitting)",
    font_size=14)

add_textbox(sl, 7.2, 4.8, 5.3, 0.5, "What the Model Learns", font_size=20, color=ORANGE, bold=True)
add_bullet_slide(sl, 7.2, 5.3, 5.3, 1.8, [
    "Mapping from mesh geometry + BCs \u2192 stress field",
    "Implicitly learns material response without PDEs",
    "Trained model runs in ~5 ms vs. ~30 s for ANSYS",
], font_size=16)

slide_number(sl, 9, total_slides)

# ═══════════════════════════════════════════════════════════════
# SLIDE 10 — GNN vs MLP Comparison
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(sl, BG_DARK)
section_bar(sl)
add_textbox(sl, 0.8, 0.4, 10, 0.7, "Two Surrogate Approaches: MLP vs. GNN", font_size=32, color=ACCENT, bold=True)

# MLP column
shape = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.5), Inches(5.0))
shape.fill.solid()
shape.fill.fore_color.rgb = BG_MID
shape.line.color.rgb = ACCENT2
shape.line.width = Pt(2)
add_textbox(sl, 1.0, 1.5, 5, 0.5, "MLP (Multi-Layer Perceptron)", font_size=22, color=ACCENT2, bold=True)
add_bullet_slide(sl, 1.0, 2.2, 5, 4, [
    "Input: flat vector of 36 dimension values",
    "Output: [volume, stress, displacement]",
    "Architecture: 36 \u2192 128 \u2192 128 \u2192 3",
    "Total parameters: ~21,000",
    "Ignores mesh structure entirely",
    "Pro: very fast, simple, no mesh needed at inference",
    "Con: cannot generalize to different topologies;\n  treats geometry as a black box",
], font_size=16, color=LIGHT)

# GNN column
shape = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(1.4), Inches(5.5), Inches(5.0))
shape.fill.solid()
shape.fill.fore_color.rgb = BG_MID
shape.line.color.rgb = ACCENT
shape.line.width = Pt(2)
add_textbox(sl, 7.2, 1.5, 5, 0.5, "GNN (Graph Neural Network)", font_size=22, color=ACCENT, bold=True)
add_bullet_slide(sl, 7.2, 2.2, 5, 4, [
    "Input: full mesh graph (nodes + edges + features)",
    "Output: predicted max stress (scalar)",
    "Architecture: Encoder + 6\u00d7EdgeConv + Head + Pool",
    "Total parameters: ~400,000",
    "Exploits spatial structure of the mesh",
    "Pro: topology-aware, mirrors FEA structure;\n  physically meaningful representations",
    "Con: requires mesh generation at inference;\n  more expensive to train and evaluate",
], font_size=16, color=LIGHT)

slide_number(sl, 10, total_slides)

# ═══════════════════════════════════════════════════════════════
# SLIDE 11 — Optimization with Genetic Algorithm
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(sl, BG_DARK)
section_bar(sl)
add_textbox(sl, 0.8, 0.4, 10, 0.7, "Stage 3: Genetic Algorithm Optimization", font_size=32, color=ACCENT, bold=True)

add_textbox(sl, 0.8, 1.3, 5.5, 0.5, "Algorithm", font_size=22, color=ORANGE, bold=True)
add_bullet_slide(sl, 0.8, 1.9, 5.5, 5, [
    "Population: 200 individuals (each = 36 dimensions)",
    "Generations: 400 evolutionary cycles",
    "Selection: tournament (k=2), top 80% survive",
    "Crossover: blend \u03b1\u00b7p1 + (1-\u03b1)\u00b7p2, \u03b1 \u2208 [0.3, 0.7]",
    "Mutation: Gaussian noise, \u03c3 = 5% of range\n  per-gene probability = 16%\n  rate adapts upward after gen 100",
    "Elitism: top 4 individuals preserved unmutated",
    "Checkpointing: population saved every generation",
], font_size=16)

add_textbox(sl, 7.2, 1.3, 5.3, 0.5, "Multi-Objective Fitness", font_size=22, color=ORANGE, bold=True)
add_code_box(sl, 7.2, 1.9, 5.3, 3.0,
    "Fitness = 0.49 \u00b7 V_norm      (minimize volume)\n"
    "       + 0.49 \u00b7 S_norm      (minimize stress)\n"
    "       + 0.02 \u00b7 (1 - D_norm) (maintain diversity)\n"
    "\n"
    "Stress constraint (soft penalty):\n"
    "  penalty = max(\u03c3 - \u03c3_yield, 0) \u00b7 weight\n"
    "  added to volume score\n"
    "\n"
    "All objectives normalized to [0, 1]",
    font_size=14)

add_textbox(sl, 7.2, 5.2, 5.3, 1.5,
    "The GA uses the surrogate for evaluation: each generation evaluates "
    "200 designs in milliseconds. Without the surrogate, this would require "
    "200 ANSYS simulations per generation \u00d7 400 generations = 80,000 simulations.",
    font_size=15, color=LIGHT)

slide_number(sl, 11, total_slides)

# ═══════════════════════════════════════════════════════════════
# SLIDE 12 — Why GA + Surrogate work together
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(sl, BG_DARK)
section_bar(sl)
add_textbox(sl, 0.8, 0.4, 10, 0.7, "The Synergy: Surrogate + Evolutionary Search", font_size=32, color=ACCENT, bold=True)

add_bullet_slide(sl, 0.8, 1.4, 11, 2, [
    "The GA does not need gradients of the objective \u2192 works even though mesh generation is non-differentiable",
    "The surrogate replaces the expensive inner loop (FEA evaluation) with a fast forward pass",
    "Cost comparison: 200 pop \u00d7 400 gen = 80,000 evaluations",
], font_size=18)

# comparison table as code box
add_code_box(sl, 1.5, 3.5, 10, 2.8,
    "\u250c\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u252c\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u252c\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2510\n"
    "\u2502                       \u2502  ANSYS FEA           \u2502  GNN Surrogate       \u2502\n"
    "\u251c\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u253c\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u253c\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2524\n"
    "\u2502 Time per evaluation    \u2502  ~30 seconds         \u2502  ~5 milliseconds     \u2502\n"
    "\u2502 80,000 evaluations     \u2502  ~28 days            \u2502  ~7 minutes          \u2502\n"
    "\u2502 Requires ANSYS license \u2502  Yes                  \u2502  No                   \u2502\n"
    "\u2502 Accuracy               \u2502  Ground truth         \u2502  Approximation        \u2502\n"
    "\u2502 Mesh-aware             \u2502  Yes                  \u2502  Yes (GNN) / No (MLP) \u2502\n"
    "\u2514\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2534\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2534\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2518",
    font_size=14)

slide_number(sl, 12, total_slides)

# ═══════════════════════════════════════════════════════════════
# SLIDE 13 — Technical Stack & Project Structure
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(sl, BG_DARK)
section_bar(sl)
add_textbox(sl, 0.8, 0.4, 10, 0.7, "Technical Stack & Project Structure", font_size=32, color=ACCENT, bold=True)

add_textbox(sl, 0.8, 1.3, 5.5, 0.5, "Technologies", font_size=22, color=ORANGE, bold=True)
add_bullet_slide(sl, 0.8, 1.9, 5.5, 4, [
    "Deep Learning: PyTorch + PyTorch Geometric",
    "FEA Solver: ANSYS MAPDL 25.1 (parallel pool)",
    "CAD Engine: IRIT Geometric Modeler",
    "Optimization: Custom GA (NumPy)",
    "Visualization: PyVista",
    "Compute: HPC cluster (PBS job scripts)",
], font_size=17)

add_textbox(sl, 7.2, 1.3, 5.3, 0.5, "Project Layout", font_size=22, color=ORANGE, bold=True)
add_code_box(sl, 7.2, 1.9, 5.3, 4.8,
    "ai_fea_surrogate/\n"
    "\u251c\u2500\u2500 core/               # Mesh, Component, IRIT\n"
    "\u2502   \u251c\u2500\u2500 component.py    # Geometry \u2192 Graph conversion\n"
    "\u2502   \u251c\u2500\u2500 mesh.py         # MAPDL interface, mesh ops\n"
    "\u2502   \u2514\u2500\u2500 IritModel.py    # Parametric CAD wrapper\n"
    "\u251c\u2500\u2500 utils/              # Model definitions\n"
    "\u2502   \u251c\u2500\u2500 gnn_surrogate.py  # GNN architecture\n"
    "\u2502   \u2514\u2500\u2500 mlp_surrogate.py  # MLP architecture\n"
    "\u251c\u2500\u2500 training/           # Data gen & training\n"
    "\u251c\u2500\u2500 evaluators/         # Inference wrappers\n"
    "\u251c\u2500\u2500 optimization/       # GA + fitness\n"
    "\u2514\u2500\u2500 data/{geometry}/    # Per-geometry configs\n"
    "    \u251c\u2500\u2500 CAD_model/      # dims.json, material\n"
    "    \u251c\u2500\u2500 dataset/        # Generated .pt files\n"
    "    \u2514\u2500\u2500 checkpoints/    # Trained models",
    font_size=13)

slide_number(sl, 13, total_slides)

# ═══════════════════════════════════════════════════════════════
# SLIDE 14 — Summary & Future Work
# ═══════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(sl, BG_DARK)
section_bar(sl)
add_textbox(sl, 0.8, 0.4, 10, 0.7, "Summary & Future Directions", font_size=32, color=ACCENT, bold=True)

add_textbox(sl, 0.8, 1.3, 5.5, 0.5, "What We Built", font_size=22, color=ORANGE, bold=True)
add_bullet_slide(sl, 0.8, 1.9, 5.5, 3.5, [
    "End-to-end pipeline: parametric CAD \u2192 FEA data\n  \u2192 surrogate training \u2192 design optimization",
    "GNN surrogate that respects mesh topology",
    "~6000\u00d7 speedup over direct FEA evaluation",
    "Multi-objective GA finds low-volume, low-stress designs",
    "Modular: supports multiple geometries & evaluators",
], font_size=17)

add_textbox(sl, 7.2, 1.3, 5.3, 0.5, "Possible Future Directions", font_size=22, color=ORANGE, bold=True)
add_bullet_slide(sl, 7.2, 1.9, 5.3, 3.5, [
    "Predict full stress fields (per-node), not just max",
    "Active learning: let the GA request new FEA data\n  in regions of high uncertainty",
    "Transfer learning across geometries",
    "Mesh-free approaches (e.g., neural operators)",
    "Uncertainty quantification for surrogate predictions",
], font_size=17)

add_textbox(sl, 1, 5.8, 11, 1,
    "Core idea: leverage the structural similarity between FEA meshes and graph neural networks\n"
    "to build a fast, topology-aware surrogate that enables efficient design space exploration.",
    font_size=19, color=ACCENT, bold=True, alignment=PP_ALIGN.CENTER)

slide_number(sl, 14, total_slides)

# ── Save ──
out_path = "/home/ryonatan/projects/ai_fea_surrogate/AI_FEA_Surrogate_Presentation.pptx"
prs.save(out_path)
print(f"Saved to {out_path}")
